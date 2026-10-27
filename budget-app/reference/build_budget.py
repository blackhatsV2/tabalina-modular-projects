#!/usr/bin/env python3
"""Builds Jayrold's Oct 27, 2026 - Dec 31, 2027 budget workbook (formula-driven).

Run: python3 build_budget.py  -> output/Budget_built.xlsx (+ output/ledger.json, a Python-side
day-by-day ledger used to verify the spreadsheet formulas). See CLAUDE.md before changing anything.
"""
import datetime as dt, calendar, math, json, sys, os
from copy import copy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference, Series

D_ = dt.date
START, END = D_(2026, 10, 27), D_(2027, 12, 31)
HERE = os.path.dirname(os.path.abspath(__file__))
os.makedirs(os.path.join(HERE, "output"), exist_ok=True)
OUT = os.path.join(HERE, "output", "Budget_built.xlsx")

# ---------------- inputs ----------------
CASH, GCASH, BANK = 300, 2731, 0
SAL, THIRTEENTH, D13 = 6000.0, 0.0, D_(2026, 12, 15)
RENT, FEE = 2000.0, 53.46
FARE, WATER, MEAL = 20, 20, 100
SUN_FARE, SUN_WATER = 80, 25
SUN_ON = True  # savings defaults leave room for it; switched off (optional) after the sim
MISC = 10
BROTHER = 500
BROTHER_ON = True  # savings defaults leave room for it; switched off (optional) after the sim
BUFFER = 500

P1 = "Mon/Wed-Sat (1st Sem 2026)"
P2 = "Thu-Sat (2nd Sem 2026)"
PB = "Break: campus 2x a week"
PX = "Christmas break: no campus"
# rides Mon..Sun, meals Mon..Sun
PATTERNS = {
    P1: ([2, 0, 2, 2, 2, 4, 0], [0] * 7,
         "Mon ITE387 10am-1pm; Wed ITE297 4:30-8pm; Thu ITE298 1-6pm; Fri ITE387/ITE300 2:30-6pm; Sat 7:30am-1pm, then ITE298 ~4:30-6pm "
         "(3+ hr gap, so home for lunch = 4 rides). No bought meals."),
    P2: ([0, 0, 0, 2, 4, 2, 0], [0, 0, 0, 1, 0, 1, 0],
         "Thu 7:30am-5pm, 30-min lunch: buy meal. Fri ITE393 7:30-9:30am, then lab ~1-4:30pm (long gap, so home for lunch = 4 rides). "
         "Sat 7:30am-5:30pm, 30-min lunch: buy meal."),
    PB: ([0, 0, 0, 2, 0, 2, 0], [0] * 7,
         "My guess for breaks: 2 campus days a week (meetings, encoding, clearance, enrollment). Change to 0 if you stay home."),
    PX: ([0] * 7, [0] * 7, "No campus trips over the Christmas break."),
}
TERMS = [  # (start, name, pattern, source note)
    (D_(2026, 10, 27), "1st Sem SY 26-27: P3 exam week", P1, "Calendar invite 'PERIOD 3 (Examination Date)' Oct 26-31; grading-sheet guide: week 17 = Oct 26-Nov 1"),
    (D_(2026, 11, 1), "Sem break: grades & clearance", PB, "P3 grade encoding Nov 3-7 (calendar invite); faculty clearance advisory Oct 27"),
    (D_(2026, 11, 16), "2nd Sem SY 26-27", P2, "ASSUMED start. Not announced yet; in 2024 the 2nd sem started about Nov 18 (your CLASS SY2425 SEM2 record)"),
    (D_(2026, 12, 20), "Christmas break", PX, "ASSUMED"),
    (D_(2027, 1, 4), "2nd Sem SY 26-27 (resumes)", P2, "ASSUMED: 17 class weeks end about Mar 27, 2027"),
    (D_(2027, 3, 28), "Summer break", PB, "ASSUMED. Change the schedule if you get a summer load"),
    (D_(2027, 7, 5), "1st Sem SY 27-28 (schedule TBA)", P2, "ASSUMED: 1st sem 2026 started Jul 6 (first Monday of July). Using the Thu-Sat costs until you know your load"),
    (D_(2027, 10, 31), "Sem break: grades & clearance", PB, "ASSUMED: same rhythm as 2026"),
    (D_(2027, 11, 15), "2nd Sem SY 27-28", P2, "ASSUMED"),
    (D_(2027, 12, 19), "Christmas break", PX, "ASSUMED"),
]
HOLIDAYS = [
    (D_(2026, 11, 1), "All Saints' Day"), (D_(2026, 11, 2), "All Souls' Day"), (D_(2026, 11, 30), "Bonifacio Day"),
    (D_(2026, 12, 8), "Immaculate Conception"), (D_(2026, 12, 24), "Christmas Eve"), (D_(2026, 12, 25), "Christmas Day"),
    (D_(2026, 12, 30), "Rizal Day"), (D_(2026, 12, 31), "Last day of the year"),
    (D_(2027, 1, 1), "New Year's Day"), (D_(2027, 2, 6), "Chinese New Year"), (D_(2027, 3, 25), "Maundy Thursday"),
    (D_(2027, 3, 26), "Good Friday"), (D_(2027, 3, 27), "Black Saturday"), (D_(2027, 4, 9), "Araw ng Kagitingan"),
    (D_(2027, 5, 1), "Labor Day"), (D_(2027, 6, 12), "Independence Day"), (D_(2027, 8, 21), "Ninoy Aquino Day"),
    (D_(2027, 8, 30), "National Heroes Day"), (D_(2027, 11, 1), "All Saints' Day"), (D_(2027, 11, 2), "All Souls' Day"),
    (D_(2027, 11, 30), "Bonifacio Day"), (D_(2027, 12, 8), "Immaculate Conception"), (D_(2027, 12, 24), "Christmas Eve"),
    (D_(2027, 12, 25), "Christmas Day"), (D_(2027, 12, 30), "Rizal Day"), (D_(2027, 12, 31), "Last day of the year"),
]
HOLSET = {h[0] for h in HOLIDAYS}
SUMMER = (TERMS[5][0], TERMS[6][0] - dt.timedelta(1))

def month_iter(y0, m0, y1, m1):
    y, m = y0, m0
    while (y, m) <= (y1, m1):
        yield y, m
        m += 1
        if m > 12: y += 1; m = 1

def term_of(d):
    t = None
    for s, name, pat, _ in TERMS:
        if s <= d: t = (name, pat)
    return t

def day_cost(d):
    name, pat = term_of(d)
    wd = d.weekday()
    hol = d in HOLSET
    rides = 0 if hol else PATTERNS[pat][0][wd]
    meals = 0 if hol else PATTERNS[pat][1][wd]
    fare = rides * FARE
    water = WATER if rides > 0 else 0
    meal = meals * MEAL
    sun = (SUN_FARE + SUN_WATER) if (wd == 6 and SUN_ON) else 0
    return dict(rides=rides, fare=fare, water=water, meal=meal, sun=sun, misc=MISC, total=fare + water + meal + sun + MISC)

# ---------------- schedule events ----------------
paydays = [D_(y, m, d) for y, m in month_iter(2026, 10, 2027, 12) for d in (15, calendar.monthrange(y, m)[1])
           if START <= D_(y, m, d) <= END]
BROTHER_START = D_(2026, 11, 15)
events = []  # (date, order, type, cat, item, amount, kind, note)
for d in paydays:
    events.append((d, 0, "Income", "Salary", "Salary (payday)", SAL, ("salary", None), ""))
events.append((D13, 0, "Income", "13th Month", "13th month pay (expected)", THIRTEENTH, ("link", "$B$17"),
               "Counted as 0 until confirmed. Set the amount on Assumptions."))
for y, m in month_iter(2026, 11, 2027, 12):
    events.append((D_(y, m, 1), 1, "Bill", "Rent", "Rent (includes electricity & water)", RENT, ("link", "$B$24"), ""))
events.append((D_(2026, 11, 5), 1, "Bill", "SPayLater", "SPayLater payment", 2726.46, ("val", None), "Includes the 53.46 handling fee"))
events.append((D_(2026, 12, 5), 1, "Bill", "SPayLater", "SPayLater payment", 2672.98, ("val", None), "Handling fee not confirmed yet - see next row"))
events.append((D_(2026, 12, 5), 1, "Bill", "SPayLater", "SPayLater handling fee (possible)", FEE, ("link", "$B$25"),
               "Safety assumption. Set to 0 on Assumptions if there is no fee."))
for name, day, amts in (("Gloan 1", 6, [1029.08] * 7 + [1029.06]), ("Gloan 2", 24, [808.56] * 6 + [808.55]), ("Gloan 3", 24, [650.06] * 9)):
    y, m = 2026, 11
    for i, a in enumerate(amts):
        events.append((D_(y, m, day), 1, "Bill", "Gloan", f"{name} - payment {i+1} of {len(amts)}", a, ("val", None), "No processing fee"))
        m += 1
        if m > 12: y += 1; m = 1
for y, m in month_iter(2026, 11, 2027, 12):
    d = D_(y, m, 15)
    if d >= BROTHER_START:
        events.append((d, 2, "Bill", "Family", "Allowance for younger brother (optional)", BROTHER, ("brother", None),
                       "Optional. Counts only if Assumptions B35 = Yes. Or type an amount here on a month you choose to send."))
events.append((D_(2027, 10, 15), 2, "Bill", "Data", "Smart data reload (26GB, no expiry)", 449, ("val", None),
               "You said one 449 load lasts about a year. Move the date when it actually runs out."))
events.sort(key=lambda e: (e[0], e[1]))

def ev_amt(e):
    if e[6][0] == "salary":
        if SUMMER[0] <= e[0] <= SUMMER[1] and not PAID_SUMMER: return 0
        return SAL
    if e[6][0] == "brother":
        return BROTHER if BROTHER_ON else 0
    return e[5]
PAID_SUMMER = True

# ---------------- simulate pay periods ----------------
starts = [START] + paydays
ends = [starts[i + 1] - dt.timedelta(1) for i in range(len(starts) - 1)] + [END]
n = len(starts)
extras = [0] * n
for i, s in enumerate(starts):
    if s == D_(2026, 12, 15): extras[i] = 500  # Christmas extra
    if s == D_(2027, 12, 15): extras[i] = 500
savings_def, sim = [], []
bal = CASH + GCASH + BANK
mins = []
for i in range(n):
    a, b = starts[i], ends[i]
    days = [a + dt.timedelta(k) for k in range((b - a).days + 1)]
    pay = sum(ev_amt(e) for e in events if e[2] == "Income" and a <= e[0] <= b)
    bills = sum(ev_amt(e) for e in events if e[2] == "Bill" and a <= e[0] <= b)
    living = sum(day_cost(d)["total"] for d in days) + extras[i]
    pre = bal + pay - bills - living
    if i >= 4:  # no savings until SPayLater is done (Dec 15)
        floor_amt = 1500 if i < 6 else 1000
        if i == n - 1:  # Jan 1-14, 2028: rent + two weeks of costs + buffer
            floor_amt = 2000 + 1700 + BUFFER  # rent + ~2 weeks of Thu-Sat costs + buffer
        s = max(0, math.floor((pre - floor_amt) / 100) * 100)
    else:
        s = 0
    savings_def.append(s)
    # daily min check
    run = bal
    for d in days:
        run += sum(ev_amt(e) for e in events if e[0] == d and e[2] == "Income")
        run -= sum(ev_amt(e) for e in events if e[0] == d and e[2] == "Bill")
        run -= day_cost(d)["total"] + (extras[i] if d == a else 0) + (s if d == a else 0)
        mins.append((round(run, 2), d.isoformat()))
    bal = pre - s
    sim.append((i + 1, a.isoformat(), b.isoformat(), len(days), pay, round(bills, 2), round(living, 2), s, round(bal, 2)))
low = min(mins)
SUN_ON = False  # optional
BROTHER_ON = False  # optional: mom also supports him
print("lowest daily balance", low)
assert low[0] >= BUFFER, low

wb = Workbook()
wsS = wb.active; wsS.title = "Summary"
wsA = wb.create_sheet("Assumptions")
wsK = wb.create_sheet("School Calendar")
wsP = wb.create_sheet("Pay Periods")
wsM = wb.create_sheet("Monthly")
wsW = wb.create_sheet("Weekly")
wsD = wb.create_sheet("Daily")
wsC = wb.create_sheet("Schedule")

# ---------------- styles ----------------
FONT = "Arial"
def F(bold=False, color="000000", size=10, italic=False): return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
NAVY, MID, PALE = "1F3864", "2F5597", "D9E2F3"
BLUE, GREEN = "0000FF", "008000"
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
PESO2 = '#,##0.00;[Red]-#,##0.00;"-"'
PESO0 = '#,##0;[Red]-#,##0;"-"'
DATEF = "mmm d, yyyy"
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical="top")
GREY9 = lambda: F(False, "595959", 9)

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = F(True, NAVY, 14)
    if sub: ws["A2"] = sub; ws["A2"].font = F(False, "595959", 9, True)
def header(ws, row, labels, start_col=1, height=48):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = F(True, "FFFFFF", 9); c.fill = fill(NAVY); c.alignment = CENTER; c.border = BOX
    ws.row_dimensions[row].height = height
def band(ws, row, c1, c2, text, color=MID):
    if c2 > c1: ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    for c in range(c1, c2 + 1): ws.cell(row=row, column=c).fill = fill(color)
    cell = ws.cell(row=row, column=c1, value=text); cell.font = F(True, "FFFFFF", 9); cell.alignment = CENTER
def widths(ws, w):
    for col, x in w.items(): ws.column_dimensions[col].width = x
def inp(cell, fmt=None):
    cell.font = F(False, BLUE); cell.fill = fill("FFF2CC"); cell.border = BOX
    if fmt: cell.number_format = fmt
def sec(ws, r, text, c2=3):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c2)
    for c in range(1, c2 + 1): ws.cell(row=r, column=c).fill = fill(PALE)
    ws.cell(row=r, column=1, value=text).font = F(True, NAVY, 10)

# =====================================================================
# Assumptions
# =====================================================================
ws = wsA
title(ws, "Assumptions & inputs", "Blue text on yellow = cells you can change. Everything else recalculates from these.")
widths(ws, {"A": 46, "B": 18, "C": 92})
def arow(r, label, val, note="", fmt=None, editable=True, link=False):
    ws.cell(row=r, column=1, value=label).font = F()
    c = ws.cell(row=r, column=2, value=val)
    if editable: inp(c, fmt)
    else:
        c.font = F(False, GREEN if link else "000000"); c.border = BOX
        if fmt: c.number_format = fmt
    x = ws.cell(row=r, column=3, value=note); x.font = GREY9(); x.alignment = WRAP
sec(ws, 4, "Plan window")
arow(5, "Plan starts", START, "Today.", DATEF)
arow(6, "Plan ends", END, "End of 2027.", DATEF)
arow(7, "Week 1 starts on (Monday)", "=B5-WEEKDAY(B5,3)", "Weeks run Monday to Sunday.", DATEF, editable=False)
sec(ws, 9, "Starting money (Oct 27, 2026)")
arow(10, "Cash on hand", CASH, "From you.", PESO2)
arow(11, "GCash balance", GCASH, "From you.", PESO2)
arow(12, "Bank account", BANK, "From you.", PESO2)
arow(13, "Total starting money", "=SUM(B10:B12)", "One spendable pool.", PESO2, editable=False)
sec(ws, 15, "Income")
arow(16, "Salary per payday (planning amount)", SAL, "Paid on the 15th and at month-end, about 6,000-6,500. Planned at 6,000 to be safe.", PESO2)
arow(17, "13th month pay expected", THIRTEENTH, "0 until confirmed. Rough idea: 1/12 of basic pay earned Jul 6-Dec 2026, about 7,000-7,500. By law it's due on or before Dec 24.", PESO2)
arow(18, "13th month expected date", D13, "Only matters if you enter an amount above.", DATEF)
arow(19, "Paid during the summer break? (Yes/No)", "Yes", "Not confirmed. Ask HR whether Instructor I pay continues from April to early July without a summer load. Set to No to see what happens if it stops.")
arow(20, "Summer break starts", "='School Calendar'!B11", "From School Calendar.", DATEF, editable=False, link=True)
arow(21, "Summer break ends", "='School Calendar'!C11", "From School Calendar.", DATEF, editable=False, link=True)
sec(ws, 23, "Fixed bills")
arow(24, "Monthly rent (due on the 1st)", RENT, "Includes electricity and water (from you). First in plan: Nov 1, 2026.", PESO2)
arow(25, "SPayLater Dec 5 handling fee (possible)", FEE, "Unconfirmed, so it's included to be safe. Set to 0 if there is none.", PESO2)
sec(ws, 27, "Daily living costs (from you, except where marked)")
arow(28, "Fare per ride", FARE, "Jeepney is 15 and the e-bus is 20 anywhere in the city. Planned at 20 to be safe. The number of rides per day comes from School Calendar.", PESO0)
arow(29, "Summit water per campus day", WATER, "1 bottle every day you go to school.", PESO0)
arow(30, "Meal when lunch break is under 1 hour", MEAL, "Long breaks (about 2 hours): you go home to eat at your aunt's, which costs 2 extra rides instead.", PESO0)
arow(31, "Sunday out (optional): fare", SUN_FARE, "4 rides x 20 if you cover both of you. Lower it on Sundays you split the fare.", PESO0)
arow(32, "Sunday out (optional): 1L Summit water", SUN_WATER, "From you: fare and one 1-liter bottle only, no food.", PESO0)
arow(33, "Personal & misc per day", MISC, "MY ADDITION: about 300 a month for toiletries, haircut, small surprises. Load/data isn't needed (Wi-Fi + your 449 Smart data).", PESO0)
arow(34, "Allowance for younger brother (per month, optional)", BROTHER, "From you: 500 when you have extra. Optional, since your mom also supports him.", PESO0)
arow(35, "Plan it every month? (Yes/No)", "Yes" if BROTHER_ON else "No", "No = not counted in the plan; send it from the 'Cushion above buffer' on Pay Periods when there's room. Yes = sent every 15th from Nov 15, 2026.")
arow(36, "Plan a Sunday out every week? (Yes/No)", "Yes" if SUN_ON else "No", "No = not counted in the plan (the default). The savings still leave room for it, so switching to Yes keeps every pay period OK.")
sec(ws, 37, "Safety & goals")
arow(38, "Minimum safety buffer", BUFFER, "Money left just before each payday should not drop below this.", PESO0)
arow(39, "Emergency fund goal", 13500, "My suggestion: about 3 months of rent + living costs (3 x (2,000 + ~2,500)).", PESO0)
sec(ws, 41, "What I found in your school accounts (Gmail, Calendar, Drive)")
found = [
    "1st Sem SY 26-27 began Jul 6, 2026 (17-week calendar in your Global Courses grading sheets). P3 exams run Oct 26-31, 2026 (calendar invite). P3 grade encoding: Nov 3-7.",
    "Registrar's faculty clearance advisory came out Oct 27, 2026. HR's Oct 26 email: your END-OF-TERM PAYROLL for the 1st sem is only processed AFTER your clearance is complete. Finish it early so a November payday isn't delayed.",
    "2nd Sem SY 26-27: enrollment appraisal and final faculty loading are ongoing (Arms's Oct 26 email says your schedule has 'minor changes'). No official start date found, so I'm assuming Nov 16, 2026 (in 2024 it started about Nov 18).",
    "Christmas break, end of 2nd sem, summer break and SY 27-28 dates are my estimates. Update them on School Calendar when the school announces them; every sheet follows.",
    "Holidays: Philippine 2027 list under Proclamation No. 1427 (Inquirer, Sep 11, 2026). Eid holidays are not yet dated. Iloilo-only days (e.g. Dinagyang) are not included.",
]
for i, t in enumerate(found):
    r = 42 + i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(row=r, column=1, value=f"{i+1}. {t}"); c.font = F(size=9); c.alignment = WRAP; ws.row_dimensions[r].height = 26
sec(ws, 48, "Other assumptions")
notes = [
    "Paydays are the 15th and the last day of each month, 6,000 each. Rent 2,000 on the 1st. SPayLater and the 3 Gloans exactly as you listed.",
    "You eat at your aunt's when you're home, so home days cost nothing except the small daily misc amount.",
    "Rides: 1 ride to school and 1 ride home. Days with a long midday gap add 2 rides (home for lunch and back). No rides on holidays.",
    "The SY 27-28 terms use the Thu-Sat costs (the higher of your two schedules: school-day costs about 420 a week vs 340) until you know your load.",
    "Each pay period's 'Savings set-aside' is money you move out of spending (e.g. to a bank or GSave) on payday.",
    "A Christmas extra of 500 is planned for the Dec 15 pay periods (2026 and 2027). Change it on Pay Periods.",
    "The last payday (Dec 31, 2027) keeps enough spendable money for Jan 1, 2028 rent and the first two weeks of January.",
]
for i, t in enumerate(notes):
    r = 49 + i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(row=r, column=1, value=f"{i+1}. {t}"); c.font = F(size=9); c.alignment = WRAP; ws.row_dimensions[r].height = 22
dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0"); ws.add_data_validation(dv)
for ref in ("B10", "B11", "B12", "B16", "B17", "B24", "B25", "B28", "B29", "B30", "B31", "B32", "B33", "B34", "B38", "B39"): dv.add(ref)
dvy = DataValidation(type="list", formula1='"Yes,No"'); ws.add_data_validation(dvy); dvy.add("B19"); dvy.add("B35"); dvy.add("B36")

# =====================================================================
# School Calendar
# =====================================================================
ws = wsK
title(ws, "School calendar & what each day costs", "Terms decide which weekly schedule applies. The schedule decides rides and bought meals per weekday. "
      "Holidays mean no trip. Edit the blue cells when PHINMA-UI announces dates or your 2nd sem load changes.")
widths(ws, {"A": 30, "B": 13, "C": 13, "D": 32, "E": 28, "F": 8, "G": 8, "H": 8, "I": 70})
band(ws, 4, 1, 9, "Terms (one row starts where the previous ends)")
header(ws, 5, ["#", "Starts", "Ends", "Term / period", "Weekly schedule used", "", "", "", "Where the date comes from"], height=30)
ws.merge_cells("E5:H5")
T0 = 6
for i, (s, name, pat, src) in enumerate(TERMS):
    r = T0 + i
    ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center")
    b = ws.cell(row=r, column=2)
    if i == 0: b.value = "=Assumptions!$B$5"; b.font = F(False, GREEN); b.number_format = DATEF; b.border = BOX
    else: b.value = s; inp(b, DATEF)
    c = ws.cell(row=r, column=3, value=("=Assumptions!$B$6" if i == len(TERMS) - 1 else f"=B{r+1}-1")); c.number_format = DATEF
    ws.cell(row=r, column=4, value=name); inp(ws.cell(row=r, column=4))
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    e = ws.cell(row=r, column=5, value=pat); inp(e)
    x = ws.cell(row=r, column=9, value=src); x.font = GREY9(); x.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 24
    for col in (1, 3): ws.cell(row=r, column=col).font = F(); ws.cell(row=r, column=col).border = BOX
TL = T0 + len(TERMS) - 1
TS, TN, TP = f"'School Calendar'!$B${T0}:$B${TL}", f"'School Calendar'!$D${T0}:$D${TL}", f"'School Calendar'!$E${T0}:$E${TL}"

R0 = 19
band(ws, R0 - 1, 1, 9, "Rides per day by weekly schedule (1 ride = one jeepney / e-bus trip)")
header(ws, R0, ["Weekly schedule", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun", "How I read your class schedule"], height=24)
M0 = R0 + 8
band(ws, M0 - 1, 1, 9, "Bought meals per day (lunch break under 1 hour)")
header(ws, M0, ["Weekly schedule", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun", "Weekly cost of this schedule (₱)"], height=24)
pat_names = list(PATTERNS)
for i, p in enumerate(pat_names):
    rr, mr = R0 + 1 + i, M0 + 1 + i
    ws.cell(row=rr, column=1, value=p).font = F(True)
    ws.cell(row=mr, column=1, value=f"=A{rr}").font = F(True, GREEN)
    for k in range(7):
        inp(ws.cell(row=rr, column=2 + k, value=PATTERNS[p][0][k]), "0"); ws.cell(row=rr, column=2 + k).alignment = Alignment(horizontal="center")
        inp(ws.cell(row=mr, column=2 + k, value=PATTERNS[p][1][k]), "0"); ws.cell(row=mr, column=2 + k).alignment = Alignment(horizontal="center")
    x = ws.cell(row=rr, column=9, value=PATTERNS[p][2]); x.font = GREY9(); x.alignment = WRAP
    ws.row_dimensions[rr].height = 38
    wk = ws.cell(row=mr, column=9, value=(f"=SUM(B{rr}:H{rr})*Assumptions!$B$28+COUNTIF(B{rr}:H{rr},\">0\")*Assumptions!$B$29"
                                          f"+SUM(B{mr}:H{mr})*Assumptions!$B$30+IF(Assumptions!$B$36=\"Yes\",Assumptions!$B$31+Assumptions!$B$32,0)+7*Assumptions!$B$33"))
    wk.number_format = PESO0; wk.font = F(True); wk.alignment = Alignment(horizontal="left")
    for col in range(1, 10):
        ws.cell(row=rr, column=col).border = BOX; ws.cell(row=mr, column=col).border = BOX
RN, RG = f"'School Calendar'!$A${R0+1}:$A${R0+len(pat_names)}", f"'School Calendar'!$B${R0+1}:$H${R0+len(pat_names)}"
MN, MG = f"'School Calendar'!$A${M0+1}:$A${M0+len(pat_names)}", f"'School Calendar'!$B${M0+1}:$H${M0+len(pat_names)}"
ws.cell(row=M0 + len(pat_names) + 1, column=9, value="Weekly cost = fares + water + bought meals + Sunday out (only if Assumptions B36 = Yes) + 7 days of misc.").font = F(False, "595959", 8, True)
dvp = DataValidation(type="list", formula1=f"=$A${R0+1}:$A${R0+len(pat_names)}"); ws.add_data_validation(dvp); dvp.add(f"E{T0}:E{TL}")

H0 = M0 + len(pat_names) + 4
band(ws, H0 - 1, 1, 4, "Holidays (no trip to campus that day)")
header(ws, H0, ["Holiday", "Date", "Day", "Source"], height=24)
for i, (d, nm) in enumerate(HOLIDAYS):
    r = H0 + 1 + i
    ws.cell(row=r, column=1, value=nm).font = F()
    inp(ws.cell(row=r, column=2, value=d), DATEF)
    c = ws.cell(row=r, column=3, value=f"=B{r}"); c.number_format = "ddd"; c.font = F()
    ws.cell(row=r, column=4, value=("Proclamation No. 1427 (2027 list)" if d.year == 2027 else "2026 national holiday list")).font = GREY9()
    for col in range(1, 5): ws.cell(row=r, column=col).border = BOX
HL = H0 + len(HOLIDAYS) + 10
for r in range(H0 + len(HOLIDAYS) + 1, HL + 1):
    for col in range(1, 5): ws.cell(row=r, column=col).border = BOX
    ws.cell(row=r, column=2).number_format = DATEF
HOL = f"'School Calendar'!$B${H0+1}:$B${HL}"
ws.freeze_panes = "A4"

# =====================================================================
# Schedule
# =====================================================================
ws = wsC
title(ws, "Schedule of paydays and bills", "Add rows at the bottom (any order). Every sheet picks them up. Types: Income or Bill. Mark Paid? = Yes once settled.")
header(ws, 4, ["Date", "Type", "Category", "Item", "Amount (₱)", "Paid?", "Note"], height=26)
widths(ws, {"A": 14, "B": 9, "C": 12, "D": 38, "E": 14, "F": 8, "G": 66})
for i, (d, _o, typ, cat, item, amt, kind, note) in enumerate(events):
    r = 5 + i
    c = ws.cell(row=r, column=1, value=d); c.number_format = DATEF; c.font = F()
    if cat == "13th Month": c.value = "=Assumptions!$B$18"; c.font = F(False, GREEN)
    for col, v in ((2, typ), (3, cat), (4, item)): ws.cell(row=r, column=col, value=v).font = F()
    a = ws.cell(row=r, column=5); a.number_format = PESO2
    if kind[0] == "link": a.value = f"=Assumptions!{kind[1]}"; a.font = F(False, GREEN)
    elif kind[0] == "brother":
        a.value = '=IF(Assumptions!$B$35="Yes",Assumptions!$B$34,0)'; a.font = F()
    elif kind[0] == "salary":
        a.value = f'=IF(AND(Assumptions!$B$19="No",A{r}>=Assumptions!$B$20,A{r}<=Assumptions!$B$21),0,Assumptions!$B$16)'; a.font = F()
    else: a.value = amt; a.font = F(False, BLUE)
    ws.cell(row=r, column=6).font = F()
    ws.cell(row=r, column=7, value=note).font = GREY9()
    for col in range(1, 8): ws.cell(row=r, column=col).border = BOX
for r in range(5 + len(events), 301):
    for col in range(1, 8): ws.cell(row=r, column=col).border = BOX
    ws.cell(row=r, column=1).number_format = DATEF; ws.cell(row=r, column=5).number_format = PESO2
ws.freeze_panes = "A5"
x = DataValidation(type="list", formula1='"Income,Bill"'); ws.add_data_validation(x); x.add("B5:B300")
x = DataValidation(type="list", formula1='"Yes"'); ws.add_data_validation(x); x.add("F5:F300")
ws.conditional_formatting.add("A5:G300", FormulaRule(formula=['$F5="Yes"'], font=Font(name=FONT, color="999999", strike=True)))
ws.conditional_formatting.add("A5:G300", FormulaRule(formula=['$B5="Income"'], fill=fill("E2EFDA")))
SD, SB, SC, SDI, SE = ("Schedule!$A$5:$A$300", "Schedule!$B$5:$B$300", "Schedule!$C$5:$C$300", "Schedule!$D$5:$D$300", "Schedule!$E$5:$E$300")

# =====================================================================
# Pay Periods
# =====================================================================
ws = wsP
title(ws, "Pay periods: the control panel", "One row per payday. Living costs come from the Daily sheet (fares, water, meals, Sundays, misc). Change the blue cells "
      "(savings set-aside, extra spending). 'Money at period end' is what's left just before the next payday: keep it above your safety buffer.")
header(ws, 5, ["#", "Starts (payday)", "Ends", "Days", "Pay in (₱)", "Bills due (₱)", "Money at start (₱)", "Savings set-aside (₱)",
               "Extra spending (₱)", "Campus days", "Living costs incl. extra (₱)", "Money at period end (₱)", "Cushion above buffer (₱)",
               "Status", "Savings balance (₱)", "Heads-up"])
widths(ws, {"A": 5, "B": 14, "C": 13, "D": 6, "E": 11, "F": 12, "G": 13, "H": 13, "I": 11, "J": 9, "K": 13, "L": 14, "M": 13, "N": 13, "O": 13, "P": 92})
headsup = {
    D_(2026, 10, 27): "P3 exam week (Wed-Sat). You start with 3,031 (cash + GCash).",
    D_(2026, 10, 31): "Heaviest bills: rent (Nov 1), SPayLater (Nov 5), Gloan 1 (Nov 6). Sem break: grade encoding Nov 3-7. Finish your faculty clearance: HR holds end-of-term payroll until it's done.",
    D_(2026, 11, 15): "Gloans 2 & 3 start (Nov 24). 2nd sem assumed to start Nov 16 (Thu-Sat). Don't spend the leftover: the next period is heavy again.",
    D_(2026, 11, 30): "Rent, last SPayLater (Dec 5, plus possible fee) and Gloan 1. Pay barely covers the bills; carry-over pays for daily costs.",
    D_(2026, 12, 15): "SPayLater done. Christmas extra of 500 planned. If a 13th month comes, enter it on Assumptions and raise this period's savings.",
    D_(2026, 12, 31): "First period of 2027: rent + Gloan 1. Classes assumed to resume Jan 4.",
    D_(2027, 3, 15): "Holy Week Mar 25-27 (no classes). 2nd sem assumed to end about Mar 27.",
    D_(2027, 3, 31): "Summer break starts (assumed). Check with HR that pay continues (Assumptions B19).",
    D_(2027, 5, 15): "Gloan 2 final payment (May 24).",
    D_(2027, 5, 31): "Gloan 1 final payment (Jun 6).",
    D_(2027, 6, 30): "1st sem SY 27-28 assumed to start Jul 5. Update School Calendar with your new load.",
    D_(2027, 7, 15): "Gloan 3 final payment (Jul 24): your last loan.",
    D_(2027, 7, 31): "All loans cleared. Only rent is fixed from here, so savings jump.",
    D_(2027, 10, 15): "Smart data reload (449) planned this period.",
    D_(2027, 12, 15): "Christmas extra of 500 planned.",
    D_(2027, 12, 31): "Last day of the plan. This payday keeps enough for Jan 1, 2028 rent + early January, so only the rest goes to savings.",
}
P0 = 6; PL = P0 + n - 1
DR = None  # defined below after Daily layout constants
D0 = 6; ndays = (END - START).days + 1; DL = D0 + ndays - 1
DR = lambda col: f"Daily!${col}${D0}:${col}${DL}"
for i in range(n):
    r = P0 + i
    ws.cell(row=r, column=1, value=i + 1)
    b = ws.cell(row=r, column=2, value=("=Assumptions!$B$5" if i == 0 else starts[i])); b.number_format = DATEF
    c = ws.cell(row=r, column=3, value=("=Assumptions!$B$6" if i == n - 1 else f"=B{r+1}-1")); c.number_format = DATEF
    ws.cell(row=r, column=4, value=f"=C{r}-B{r}+1")
    ws.cell(row=r, column=5, value=f'=SUMIFS({SE},{SD},">="&B{r},{SD},"<="&C{r},{SB},"Income")')
    ws.cell(row=r, column=6, value=f'=SUMIFS({SE},{SD},">="&B{r},{SD},"<="&C{r},{SB},"Bill")')
    ws.cell(row=r, column=7, value=("=Assumptions!$B$13" if i == 0 else f"=L{r-1}"))
    inp(ws.cell(row=r, column=8, value=savings_def[i]), PESO0)
    inp(ws.cell(row=r, column=9, value=extras[i]), PESO0)
    ws.cell(row=r, column=10, value=f'=COUNTIFS({DR("D")},A{r},{DR("H")},">0")')
    ws.cell(row=r, column=11, value=f'=SUMIFS({DR("O")},{DR("D")},A{r})')
    ws.cell(row=r, column=12, value=f"=G{r}+E{r}-F{r}-H{r}-K{r}")
    ws.cell(row=r, column=13, value=f"=L{r}-Assumptions!$B$38")
    ws.cell(row=r, column=14, value=f'=IF(L{r}<0,"SHORTFALL",IF(L{r}<Assumptions!$B$38,"Below buffer","OK"))')
    ws.cell(row=r, column=15, value=(f"=H{r}" if i == 0 else f"=O{r-1}+H{r}"))
    hc = ws.cell(row=r, column=16, value=headsup.get(starts[i], "")); hc.font = GREY9(); hc.alignment = Alignment(wrap_text=True, vertical="center")
    if len(headsup.get(starts[i], "")) > 105: ws.row_dimensions[r].height = 26
    for col in range(1, 17):
        cell = ws.cell(row=r, column=col); cell.border = BOX
        if col not in (8, 9, 16):
            cell.font = F(False, GREEN) if (col == 7 and i == 0) or (col == 2 and i == 0) or (col == 3 and i == n - 1) else F()
        if col in (5, 6, 7, 11, 12, 13, 15): cell.number_format = PESO2
        if col in (1, 4, 10, 14): cell.alignment = Alignment(horizontal="center")
ws.conditional_formatting.add(f"N{P0}:N{PL}", FormulaRule(formula=[f'$N{P0}="SHORTFALL"'], fill=fill("F8CBAD"), font=Font(name=FONT, bold=True, color="9C0006")))
ws.conditional_formatting.add(f"N{P0}:N{PL}", FormulaRule(formula=[f'$N{P0}="Below buffer"'], fill=fill("FFE699")))
ws.conditional_formatting.add(f"N{P0}:N{PL}", FormulaRule(formula=[f'$N{P0}="OK"'], fill=fill("E2EFDA"), font=Font(name=FONT, color="375623")))
ws.freeze_panes = "C6"
x = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0"); ws.add_data_validation(x); x.add(f"H{P0}:I{PL}")
ws.cell(row=PL + 2, column=1, value="Money carries forward: if you raise savings or extra spending in one period, check that Status stays OK for the next periods too.").font = F(False, "595959", 9, True)
PPS, PPH, PPI = f"'Pay Periods'!$B${P0}:$B${PL}", f"'Pay Periods'!$H${P0}:$H${PL}", f"'Pay Periods'!$I${P0}:$I${PL}"

# =====================================================================
# Daily
# =====================================================================
ws = wsD
title(ws, "Daily plan", "One row per day. Rides, water and meals follow the School Calendar. Green = payday, amber = bill due, blue = today. "
      "Log what you actually spent (not counting bills) in the blue 'Actual spent' column.")
ws["A3"] = "Safety buffer (₱):"; ws["A3"].font = F(True, size=9)
ws["B3"] = "=Assumptions!$B$38"; ws["B3"].font = F(False, GREEN, 9); ws["B3"].number_format = PESO0
band(ws, 4, 1, 4, "When"); band(ws, 4, 5, 8, "School day"); band(ws, 4, 9, 15, "Living costs (planned)")
band(ws, 4, 16, 19, "Money in / bills / savings"); band(ws, 4, 20, 21, "Balances"); band(ws, 4, 22, 24, "Your tracking", "BF9000")
header(ws, 5, ["Date", "Day", "Week #", "Pay period #", "Term / period", "Weekly schedule", "Day type", "Rides",
               "Fare (₱)", "Water (₱)", "Meal at school (₱)", "Sunday out (₱)", "Misc (₱)", "Extra (₱)", "Living total (₱)",
               "Income (₱)", "Bills due (₱)", "What's due / paid today", "Moved to savings (₱)",
               "Planned balance, end of day (₱)", "Savings balance (₱)", "Actual spent (₱)", "Under / (over) (₱)", "Running under / (over) (₱)"], height=52)
widths(ws, {"A": 13, "B": 6, "C": 7, "D": 7, "E": 27, "F": 22, "G": 11, "H": 6, "I": 8, "J": 8, "K": 9, "L": 9, "M": 7, "N": 8, "O": 10,
            "P": 10, "Q": 10, "R": 34, "S": 10, "T": 13, "U": 11, "V": 10, "W": 10, "X": 11})
for i in range(ndays):
    r = D0 + i
    a = ws.cell(row=r, column=1, value=("=Assumptions!$B$5" if i == 0 else f"=A{r-1}+1")); a.number_format = DATEF
    ws.cell(row=r, column=2, value=f"=A{r}").number_format = "ddd"
    ws.cell(row=r, column=3, value=f"=INT((A{r}-Assumptions!$B$7)/7)+1")
    ws.cell(row=r, column=4, value=f"=MATCH(A{r},{PPS},1)")
    ws.cell(row=r, column=5, value=f"=INDEX({TN},MATCH(A{r},{TS},1))")
    ws.cell(row=r, column=6, value=f"=INDEX({TP},MATCH(A{r},{TS},1))")
    ws.cell(row=r, column=7, value=f'=IF(COUNTIF({HOL},A{r})>0,"Holiday",IF(H{r}>0,"Campus day",IF(AND(WEEKDAY(A{r},2)=7,Assumptions!$B$36="Yes"),"Sunday out","Home")))')
    ws.cell(row=r, column=8, value=f"=IF(COUNTIF({HOL},A{r})>0,0,INDEX({RG},MATCH(F{r},{RN},0),WEEKDAY(A{r},2)))")
    ws.cell(row=r, column=9, value=f"=H{r}*Assumptions!$B$28")
    ws.cell(row=r, column=10, value=f"=IF(H{r}>0,Assumptions!$B$29,0)")
    ws.cell(row=r, column=11, value=f"=IF(COUNTIF({HOL},A{r})>0,0,INDEX({MG},MATCH(F{r},{MN},0),WEEKDAY(A{r},2)))*Assumptions!$B$30")
    ws.cell(row=r, column=12, value=f'=IF(AND(WEEKDAY(A{r},2)=7,Assumptions!$B$36="Yes"),Assumptions!$B$31+Assumptions!$B$32,0)')
    ws.cell(row=r, column=13, value="=Assumptions!$B$33")
    ws.cell(row=r, column=14, value=f"=IF(A{r}=INDEX({PPS},D{r}),INDEX({PPI},D{r}),0)")
    ws.cell(row=r, column=15, value=f"=SUM(I{r}:N{r})")
    ws.cell(row=r, column=16, value=f'=SUMIFS({SE},{SD},$A{r},{SB},"Income")')
    ws.cell(row=r, column=17, value=f'=SUMIFS({SE},{SD},$A{r},{SB},"Bill")')
    ws.cell(row=r, column=18, value=(f'=IF(COUNTIF({SD},A{r})=0,"",INDEX({SDI},MATCH(A{r},{SD},0))'
                                     f'&IF(COUNTIF({SD},A{r})>1," +"&(COUNTIF({SD},A{r})-1)&" more",""))'))
    ws.cell(row=r, column=19, value=f"=IF(A{r}=INDEX({PPS},D{r}),INDEX({PPH},D{r}),0)")
    prev = "Assumptions!$B$13" if i == 0 else f"T{r-1}"
    ws.cell(row=r, column=20, value=f"={prev}+P{r}-Q{r}-O{r}-S{r}")
    ws.cell(row=r, column=21, value=(f"=S{r}" if i == 0 else f"=U{r-1}+S{r}"))
    inp(ws.cell(row=r, column=22), PESO2)
    ws.cell(row=r, column=23, value=f'=IF(V{r}="","",O{r}-V{r})')
    ws.cell(row=r, column=24, value=f'=IF(V{r}="","",SUM($W${D0}:W{r}))')
    for col in range(1, 25):
        c = ws.cell(row=r, column=col); c.border = BOX
        if col != 22: c.font = F(size=9) if col in (5, 6, 7, 18) else F()
        if col in (16, 17, 20, 21, 23, 24): c.number_format = PESO2
        if col in (9, 10, 11, 12, 13, 14, 15, 19): c.number_format = PESO0
        if col in (2, 3, 4, 7, 8): c.alignment = Alignment(horizontal="center")
ws.freeze_panes = "C6"
ws.auto_filter.ref = f"A5:X{DL}"
rng = f"A{D0}:H{DL}"
ws.conditional_formatting.add(f"A{D0}:X{DL}", FormulaRule(formula=[f"$A{D0}=TODAY()"], fill=fill("BDD7EE"), font=Font(name=FONT, bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f"$P{D0}>0"], fill=fill("E2EFDA")))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f"$Q{D0}>0"], fill=fill("FFF2CC")))
ws.conditional_formatting.add(f"G{D0}:G{DL}", FormulaRule(formula=[f'$G{D0}="Holiday"'], fill=fill("E4DFEC"), font=Font(name=FONT, color="5B2C6F")))
ws.conditional_formatting.add(f"G{D0}:G{DL}", FormulaRule(formula=[f'$G{D0}="Campus day"'], font=Font(name=FONT, bold=True, color="1F3864")))
ws.conditional_formatting.add(f"T{D0}:T{DL}", FormulaRule(formula=[f"$T{D0}<0"], fill=fill("F8CBAD"), font=Font(name=FONT, bold=True, color="9C0006")))
ws.conditional_formatting.add(f"T{D0}:T{DL}", FormulaRule(formula=[f"$T{D0}<$B$3"], fill=fill("FFE699")))
ws.conditional_formatting.add(f"A{D0}:X{DL}", FormulaRule(formula=[f"DAY($A{D0})=1"], border=Border(top=Side(style="medium", color="1F3864"))))

# =====================================================================
# Weekly
# =====================================================================
ws = wsW
w1 = START - dt.timedelta(days=START.weekday())
nweeks = ((END - w1).days // 7) + 1
W0, WL = 6, 6 + nweeks - 1
title(ws, "Weekly plan (Monday to Sunday)", "Roll-up of the Daily sheet. Week 1 and the last week are partial. Actual columns fill in as you log on Daily.")
band(ws, 4, 1, 6, "Week"); band(ws, 4, 7, 8, "Money in / bills"); band(ws, 4, 9, 15, "Living costs (planned)")
band(ws, 4, 16, 19, "Savings & balances"); band(ws, 4, 20, 21, "Your tracking", "BF9000")
header(ws, 5, ["Week #", "Week starts (Mon)", "Week ends (Sun)", "Days in plan", "Term / period", "Campus days", "Income (₱)", "Bills due (₱)",
               "Fares (₱)", "Water (₱)", "Meals at school (₱)", "Sunday out (₱)", "Misc (₱)", "Extra (₱)", "Living total (₱)",
               "Moved to savings (₱)", "Net change (₱)", "Planned balance, end of week (₱)", "Savings balance (₱)", "Actual spent (₱)", "Under / (over) (₱)"], height=52)
widths(ws, {"A": 7, "B": 13, "C": 13, "D": 7, "E": 28, "F": 8, "G": 11, "H": 11, "I": 9, "J": 9, "K": 10, "L": 10, "M": 8, "N": 8, "O": 11,
            "P": 11, "Q": 11, "R": 14, "S": 12, "T": 11, "U": 11})
for i in range(nweeks):
    r = W0 + i
    ws.cell(row=r, column=1, value=i + 1)
    ws.cell(row=r, column=2, value=f"=Assumptions!$B$7+(A{r}-1)*7").number_format = DATEF
    ws.cell(row=r, column=3, value=f"=B{r}+6").number_format = DATEF
    ws.cell(row=r, column=4, value=f"=MAX(0,MIN(C{r},Assumptions!$B$6)-MAX(B{r},Assumptions!$B$5)+1)")
    ws.cell(row=r, column=5, value=f"=INDEX({DR('E')},MATCH(MAX(B{r},Assumptions!$B$5),{DR('A')},0))")
    ws.cell(row=r, column=6, value=f'=COUNTIFS({DR("C")},$A{r},{DR("H")},">0")')
    for col, src in ((7, "P"), (8, "Q"), (9, "I"), (10, "J"), (11, "K"), (12, "L"), (13, "M"), (14, "N"), (15, "O"), (16, "S")):
        ws.cell(row=r, column=col, value=f"=SUMIFS({DR(src)},{DR('C')},$A{r})")
    ws.cell(row=r, column=17, value=f"=G{r}-H{r}-O{r}-P{r}")
    ws.cell(row=r, column=18, value=f"=INDEX({DR('T')},MATCH(MIN($C{r},Assumptions!$B$6),{DR('A')},0))")
    ws.cell(row=r, column=19, value=f"=INDEX({DR('U')},MATCH(MIN($C{r},Assumptions!$B$6),{DR('A')},0))")
    ws.cell(row=r, column=20, value=f'=IF(COUNTIFS({DR("C")},$A{r},{DR("V")},"<>")=0,"",SUMIFS({DR("V")},{DR("C")},$A{r}))')
    ws.cell(row=r, column=21, value=f'=IF(T{r}="","",SUMIFS({DR("W")},{DR("C")},$A{r}))')
    for col in range(1, 22):
        c = ws.cell(row=r, column=col); c.font = F(size=9) if col == 5 else F(); c.border = BOX
        if col in (1, 4, 6): c.alignment = Alignment(horizontal="center")
        if col in (7, 8, 17, 18, 19, 20, 21): c.number_format = PESO2
        if col in (9, 10, 11, 12, 13, 14, 15, 16): c.number_format = PESO0
ws.freeze_panes = "C6"
ws.conditional_formatting.add(f"A{W0}:U{WL}", FormulaRule(formula=[f"AND(TODAY()>=$B{W0},TODAY()<=$C{W0})"], fill=fill("BDD7EE"), font=Font(name=FONT, bold=True)))
ws.conditional_formatting.add(f"R{W0}:R{WL}", FormulaRule(formula=[f"$R{W0}<0"], fill=fill("F8CBAD")))
tr = WL + 1
ws.cell(row=tr, column=1, value="Total")
for col in range(1, 22):
    L_ = ws.cell(row=tr, column=col).column_letter; c = ws.cell(row=tr, column=col)
    if col in (4, 6) + tuple(range(7, 18)) + (20, 21): c.value = f"=SUM({L_}{W0}:{L_}{WL})"
    c.font = F(True); c.fill = fill(PALE); c.border = BOX
    c.number_format = PESO2 if col in (7, 8, 17, 20, 21) else (PESO0 if col >= 9 else "0")

# =====================================================================
# Monthly
# =====================================================================
ws = wsM
months = list(month_iter(2026, 10, 2027, 12))
M0r, MLr = 6, 6 + len(months) - 1
title(ws, "Monthly plan", "Roll-up of the Daily sheet. Oct 2026 is just Oct 27-31. 'Bills % of income' shows how squeezed each month is.")
band(ws, 4, 1, 5, "Month"); band(ws, 4, 6, 6, "In"); band(ws, 4, 7, 12, "Bills & sending (Schedule)"); band(ws, 4, 13, 19, "Living costs (planned)")
band(ws, 4, 20, 21, "Savings & net"); band(ws, 4, 22, 24, "Money at month-end"); band(ws, 4, 25, 25, "Ratio"); band(ws, 4, 26, 27, "Your tracking", "BF9000")
header(ws, 5, ["Month", "From", "To", "Days", "Campus days", "Income (₱)", "Rent (₱)", "SPayLater (₱)", "Gloan (₱)", "Brother (₱)", "Other bills (₱)",
               "Total bills (₱)", "Fares (₱)", "Water (₱)", "Meals at school (₱)", "Sunday out (₱)", "Misc (₱)", "Extra (₱)", "Living total (₱)",
               "Moved to savings (₱)", "Net change in spendable (₱)", "Spendable balance (₱)", "Savings balance (₱)", "Total money (₱)",
               "Bills % of income", "Actual spent (₱)", "Under / (over) (₱)"], height=52)
widths(ws, {"A": 17, "B": 12, "C": 12, "D": 6, "E": 8, "F": 11, "G": 9, "H": 11, "I": 10, "J": 9, "K": 10, "L": 11, "M": 9, "N": 8, "O": 10,
            "P": 10, "Q": 8, "R": 8, "S": 11, "T": 11, "U": 12, "V": 12, "W": 12, "X": 12, "Y": 9, "Z": 11, "AA": 11})
for i, (y, m) in enumerate(months):
    r = M0r + i
    first = D_(y, m, 1); lastd = D_(y, m, calendar.monthrange(y, m)[1])
    frm, to = max(first, START), min(lastd, END)
    ws.cell(row=r, column=1, value=first.strftime("%b %Y") + (" (27-31)" if frm != first else ""))
    ws.cell(row=r, column=2, value=frm).number_format = DATEF
    ws.cell(row=r, column=3, value=to).number_format = DATEF
    ws.cell(row=r, column=4, value=f"=C{r}-B{r}+1")
    dbase = f'{DR("A")},">="&$B{r},{DR("A")},"<="&$C{r}'
    ws.cell(row=r, column=5, value=f'=COUNTIFS({dbase},{DR("H")},">0")')
    base = f'{SD},">="&$B{r},{SD},"<="&$C{r}'
    ws.cell(row=r, column=6, value=f'=SUMIFS({SE},{base},{SB},"Income")')
    for col, cat in ((7, "Rent"), (8, "SPayLater"), (9, "Gloan"), (10, "Family")):
        ws.cell(row=r, column=col, value=f'=SUMIFS({SE},{base},{SC},"{cat}")')
    ws.cell(row=r, column=11, value=f"=L{r}-G{r}-H{r}-I{r}-J{r}")
    ws.cell(row=r, column=12, value=f'=SUMIFS({SE},{base},{SB},"Bill")')
    for col, src in ((13, "I"), (14, "J"), (15, "K"), (16, "L"), (17, "M"), (18, "N"), (19, "O"), (20, "S")):
        ws.cell(row=r, column=col, value=f"=SUMIFS({DR(src)},{dbase})")
    ws.cell(row=r, column=21, value=f"=F{r}-L{r}-S{r}-T{r}")
    ws.cell(row=r, column=22, value=f"=INDEX({DR('T')},MATCH($C{r},{DR('A')},0))")
    ws.cell(row=r, column=23, value=f"=INDEX({DR('U')},MATCH($C{r},{DR('A')},0))")
    ws.cell(row=r, column=24, value=f"=V{r}+W{r}")
    ws.cell(row=r, column=25, value=f'=IF(F{r}=0,"",L{r}/F{r})')
    ws.cell(row=r, column=26, value=f'=IF(COUNTIFS({dbase},{DR("V")},"<>")=0,"",SUMIFS({DR("V")},{dbase}))')
    ws.cell(row=r, column=27, value=f'=IF(Z{r}="","",SUMIFS({DR("W")},{dbase}))')
    for col in range(1, 28):
        c = ws.cell(row=r, column=col); c.font = F(); c.border = BOX
        if col in (4, 5): c.alignment = Alignment(horizontal="center")
        if col in (6, 7, 8, 9, 10, 11, 12, 21, 22, 23, 24, 26, 27): c.number_format = PESO2
        if col in (13, 14, 15, 16, 17, 18, 19, 20): c.number_format = PESO0
        if col == 25: c.number_format = "0%"
mt = MLr + 1
ws.cell(row=mt, column=1, value="Total / end of plan")
for col in range(1, 28):
    L_ = ws.cell(row=mt, column=col).column_letter; c = ws.cell(row=mt, column=col)
    if 4 <= col <= 21 or col in (26, 27): c.value = f"=SUM({L_}{M0r}:{L_}{MLr})"
    if col in (22, 23, 24): c.value = f"={L_}{MLr}"
    if col == 25: c.value = f'=IF(F{mt}=0,"",L{mt}/F{mt})'
    c.font = F(True); c.fill = fill(PALE); c.border = BOX
    if col in (6, 7, 8, 9, 10, 11, 12, 21, 22, 23, 24, 26, 27): c.number_format = PESO2
    if col in (13, 14, 15, 16, 17, 18, 19, 20): c.number_format = PESO0
    if col == 25: c.number_format = "0%"
    if col in (4, 5): c.alignment = Alignment(horizontal="center")
ws.freeze_panes = "B6"
ws.conditional_formatting.add(f"Y{M0r}:Y{MLr}", FormulaRule(formula=[f"AND(ISNUMBER($Y{M0r}),$Y{M0r}>=0.5)"], fill=fill("FFE699")))
ws.conditional_formatting.add(f"A{M0r}:AA{MLr}", FormulaRule(formula=[f"AND(TODAY()>=$B{M0r},TODAY()<=$C{M0r})"], fill=fill("BDD7EE")))
ws.cell(row=mt + 2, column=1, value="Net change = Income - Total bills - Living total - Moved to savings = month-to-month change in the spendable balance.").font = F(False, "595959", 9, True)
ws.cell(row=mt + 3, column=1, value="Month-end balances include that day's payday, which has to cover rent on the 1st and the early-month bills, so they look higher than what you can freely spend.").font = F(False, "595959", 9, True)

# =====================================================================
# Summary
# =====================================================================
ws = wsS
title(ws, "Budget plan: Oct 27, 2026 to Dec 31, 2027", "Philippine pesos (₱). Built from your pay dates, bills, class schedules and the PHINMA-UI calendar.")
widths(ws, {"A": 46, "B": 18, "C": 34, "D": 3})
def sband(r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for c in (1, 2, 3): ws.cell(row=r, column=c).fill = fill(NAVY)
    ws.cell(row=r, column=1, value=text).font = F(True, "FFFFFF", 10)
def srow(r, label, formula, fmt=PESO2, note=None, bold=False, green=False):
    ws.cell(row=r, column=1, value=label).font = F(bold)
    c = ws.cell(row=r, column=2, value=formula); c.number_format = fmt; c.font = F(bold, GREEN if green else "000000"); c.alignment = Alignment(horizontal="right")
    if note is not None:
        x = ws.cell(row=r, column=3, value=note); x.font = GREY9(); x.alignment = Alignment(wrap_text=True, vertical="center")
    for col in (1, 2, 3): ws.cell(row=r, column=col).border = BOX
sband(4, "Where the money goes (whole plan)")
srow(5, "Starting money (Oct 27, 2026)", "=Assumptions!$B$13", note="Cash + GCash + bank", green=True)
srow(6, "Salary income", f'=SUMIFS({SE},{SC},"Salary")', note=f'=COUNTIF({SC},"Salary")&" paydays"')
srow(7, "13th month pay counted", f'=SUMIFS({SE},{SC},"13th Month")', note="0 until you confirm an amount")
srow(8, "Total money available", "=SUM(B5:B7)", bold=True)
items = [("Rent", "G"), ("SPayLater", "H"), ("Gloan (3 loans)", "I"), ("Allowance for your brother (optional)", "J"), ("Other bills (Smart data, etc.)", "K"),
         ("Fares (jeepney / e-bus)", "M"), ("Summit water", "N"), ("Meals at school", "O"), ("Sundays out (optional)", "P"),
         ("Personal & misc", "Q"), ("Extra spending (Christmas)", "R"), ("Moved to savings", "T")]
for k, (lab, col) in enumerate(items):
    srow(9 + k, lab, f"=Monthly!{col}{mt}", green=True)
srow(21, "Left in spendable money at the end", f"=Monthly!V{MLr}", bold=True, green=True)
srow(22, "Check: available - everything above = 0", "=B8-SUM(B9:B20)-B21")
ws["C22"] = '=IF(ABS(B22)<0.005,"OK, it balances","Check the plan")'; ws["C22"].font = GREY9()

sband(24, "Typical week (fares, water, meals, misc; Sunday only if planned)")
for k, p in enumerate(pat_names):
    srow(25 + k, p, f"='School Calendar'!I{M0+1+k}", PESO0, green=True)
ws["C25"] = "Your current schedule (until Oct 31)"; ws["C26"] = "Your 2nd sem schedule"; ws["C27"] = "Sem/summer breaks (my guess: 2 campus days)"
for c in ("C25", "C26", "C27"): ws[c].font = GREY9()

sband(30, "Savings")
srow(31, "Savings balance, Dec 31, 2027", f"=Monthly!W{MLr}", bold=True, green=True)
srow(32, "Spendable + savings, Dec 31, 2027", f"=Monthly!X{MLr}", green=True)
srow(33, "Emergency fund goal", "=Assumptions!$B$39", PESO0, green=True)
srow(34, "Savings reach the goal on", f'=IFERROR(INDEX({DR("A")},COUNTIF({DR("U")},"<"&Assumptions!$B$39)+1),"Not by Dec 2027")', DATEF)

sband(36, "Watch-outs")
srow(37, "Lowest planned balance (any day)", f"=MIN({DR('T')})")
srow(38, "...on", f"=INDEX({DR('A')},MATCH(B37,{DR('T')},0))", DATEF, note="Right before payday: the tightest moment")
srow(39, "Days below the safety buffer", f'=COUNTIF({DR("T")},"<"&Assumptions!$B$38)', "0")
srow(40, "Pay periods not marked OK", f"=COUNTIF('Pay Periods'!$N${P0}:$N${PL},\"<>OK\")", "0", note="See Pay Periods")
srow(41, "Salary paid during summer break?", "=Assumptions!$B$19", "@", note="Unconfirmed: ask HR", green=True)
srow(42, "Faculty clearance", "Do it now", "@", note="HR holds end-of-term payroll until clearance is complete")

sband(44, "School calendar used (edit on School Calendar)")
for k in range(len(TERMS)):
    r = 45 + k
    ws.cell(row=r, column=1, value=f"='School Calendar'!D{T0+k}").font = F(False, GREEN)
    c = ws.cell(row=r, column=2, value=f"='School Calendar'!B{T0+k}"); c.number_format = DATEF; c.font = F(False, GREEN); c.alignment = Alignment(horizontal="right")
    c = ws.cell(row=r, column=3, value=f'="to "&TEXT(\'School Calendar\'!C{T0+k},"mmm d, yyyy")'); c.font = GREY9()
    for col in (1, 2, 3): ws.cell(row=r, column=col).border = BOX
ER = 45 + len(TERMS) + 1
sband(ER, "When the debts end")
srow(ER + 1, "SPayLater: last payment", f'=_xlfn.MAXIFS({SD},{SC},"SPayLater")', DATEF)
srow(ER + 2, "Gloan 1 (1,029.08): last payment", f'=_xlfn.MAXIFS({SD},{SDI},"Gloan 1*")', DATEF)
srow(ER + 3, "Gloan 2 (808.56): last payment", f'=_xlfn.MAXIFS({SD},{SDI},"Gloan 2*")', DATEF)
srow(ER + 4, "Gloan 3 (650.06): last payment", f'=_xlfn.MAXIFS({SD},{SDI},"Gloan 3*")', DATEF)
TR = ER + 6
sband(TR, "Today")
srow(TR + 1, "Today's date", "=TODAY()", DATEF)
srow(TR + 2, "Planned living cost today", f'=IFERROR(INDEX({DR("O")},MATCH(TODAY(),{DR("A")},0)),"outside plan")', PESO0)
srow(TR + 3, "Planned balance at end of today", f'=IFERROR(INDEX({DR("T")},MATCH(TODAY(),{DR("A")},0)),"outside plan")')
GR = TR + 5
sband(GR, "How to use this workbook")
guide = [
    "Assumptions: fare, water, meal, Sunday and brother amounts. School Calendar: term dates, rides/meals per weekday, holidays.",
    "Pay Periods is the control panel: savings set-aside and extra spending per payday, and an OK / Below buffer status.",
    "Daily: every day's rides and costs, bills due and the planned balance. Log real spending in the blue 'Actual spent' column.",
    "Weekly (Mon-Sun) and Monthly are roll-ups of Daily. Schedule lists every payday and bill; add rows at the bottom.",
    "Blue on yellow = you can edit. Black = formula. Green = pulled from another sheet.",
]
for i, t in enumerate(guide):
    r = GR + 1 + i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(row=r, column=1, value=f"{i+1}. {t}"); c.font = F(size=9); c.alignment = WRAP; ws.row_dimensions[r].height = 26

ch1 = BarChart(); ch1.type = "col"; ch1.grouping = "stacked"; ch1.overlap = 100; ch1.title = "Money at month-end (₱)"
for col, name in ((22, "Spendable"), (23, "Savings")):
    ch1.series.append(Series(Reference(wsM, min_col=col, min_row=M0r, max_row=MLr), title=name))
ch1.set_categories(Reference(wsM, min_col=1, min_row=M0r, max_row=MLr))
ch2 = BarChart(); ch2.type = "col"; ch2.grouping = "clustered"; ch2.title = "Each month: income vs bills vs living costs (₱)"
for col, name in ((6, "Income"), (12, "Bills & sending"), (19, "Living costs")):
    ch2.series.append(Series(Reference(wsM, min_col=col, min_row=M0r, max_row=MLr), title=name))
ch2.set_categories(Reference(wsM, min_col=1, min_row=M0r, max_row=MLr))
for ch, anchor in ((ch1, "E4"), (ch2, "E22")):
    ch.height = 8.5; ch.width = 22; ch.legend.position = "b"
    ch.y_axis.number_format = "#,##0"; ch.y_axis.delete = False; ch.x_axis.delete = False
    ws.add_chart(ch, anchor)

# ---- inserted into build_budget2.py: glossary sheet + hover comments on every header ----
from openpyxl.comments import Comment

G = {}  # normalized header -> (plain explanation, example / tip)
def g(keys, text, tip=""):
    for k in keys.split("|"): G[k] = (text, tip)

g("Pay period|#|Pay period #",
  "A pay period is the stretch from one payday to the day before the next payday (about 15 days). Your money has to last that whole stretch, so the plan is built around it.",
  "Example: the Oct 31 payday has to cover Oct 31 to Nov 14.")
g("Starts (payday)", "The payday that opens this pay period (the 15th or the last day of the month).")
g("Ends", "The last day of the period, the day before your next payday. This is usually the day your money is lowest.")
g("Days|Days in plan", "How many days the period, week or month covers in the plan.")
g("Pay in|Income", "Money coming in: your salary on paydays (planned at 6,000 each), plus the 13th month pay if you enter it.")
g("Bills due|Total bills",
  "Fixed payments with a due date: rent, SPayLater, Gloan, the Smart data reload, and your brother's allowance if it's switched on. These come first before any other spending.",
  "They're listed one by one on the Schedule sheet.")
g("Money at start",
  "What you have right before this period's payday arrives: the leftover from the previous period. For the first period it's your 300 cash + 2,731 GCash.")
g("Savings set-aside|Moved to savings",
  "Money you move OUT of your spending money on payday into a separate place (e.g. a bank account, GSave or Maya savings) and don't touch. This is 'pay yourself first': save on payday, not whatever is left at the end.",
  "If it stays in the same GCash wallet you'll probably spend it. Keep it somewhere separate.")
g("Extra spending|Extra",
  "A one-time planned spend for that period, on top of normal daily costs (e.g. 500 for Christmas). It's taken out on the payday that starts the period.")
g("Campus days", "Days you travel to PHINMA-UI (classes, exams, meetings). Only these days cost fare and water.")
g("Living costs incl. extra|Living total",
  "Your day-to-day spending: fares + Summit water + meals bought at school + Sunday out (if switched on) + small misc, plus any extra spending. Bills and savings are NOT included here.")
g("Money at period end",
  "What should be left the day before your next payday. Formula: Money at start + Pay in − Bills due − Savings set-aside − Living costs. It carries over as the next period's 'Money at start'.",
  "Keep this above your safety buffer (500).")
g("Cushion above buffer",
  "How much more than your 500 safety buffer you'll have at the end of the period. This is your 'free' money: you can spend it (your brother's allowance, a Sunday out) or move it to savings without breaking the plan.",
  "If it's small, be careful that period. If it's big, you can save more.")
g("Status",
  "OK = the period ends above your safety buffer. Below buffer = it ends with less than 500 (tight, but not broke). SHORTFALL = the plan runs out of money: cut spending, save less, or delay something.")
g("Savings balance",
  "The running total of everything you've moved to savings so far. Your savings pile, separate from your spending money.")
g("Heads-up", "Notes on what's special about that period: heavy bills, school dates, or things to do.")
g("Date|Day", "The calendar date / day of the week.")
g("Week #", "Week number in the plan. Weeks run Monday to Sunday; week 1 is Oct 26 to Nov 1, 2026.")
g("Term / period", "Which part of the school year the day falls in (semester, sem break, Christmas or summer break). Comes from the School Calendar sheet.")
g("Weekly schedule", "Which class schedule decides your rides and meals that day. Comes from the School Calendar sheet.")
g("Day type", "Campus day = you go to school. Home = no trip. Holiday = no classes. Sunday out = only if you switched Sunday outings on.")
g("Rides", "Number of jeepney / e-bus rides that day. 2 = to school and home. 4 = you also go home for lunch and come back.")
g("Fare|Fares", "Rides × fare per ride (planned at 20, the e-bus price, to be safe; jeepney is 15).")
g("Water", "One Summit bottle (20) on each campus day.")
g("Meal at school|Meals at school", "A 100 meal on days your lunch break is under 1 hour. On longer breaks you eat at your aunt's for free and pay extra fare instead.")
g("Sunday out", "Fare + one 1L Summit (25) on Sundays, but only if 'Plan a Sunday out every week?' is Yes on Assumptions. Otherwise 0.")
g("Misc", "A small daily amount (10) for toiletries, a haircut and small surprises. About 300 a month.")
g("What's due / paid today", "The bill or payday on that date (from Schedule). '+1 more' means there's another one that day.")
g("Planned balance, end of day|Planned balance, end of week",
  "Your spending money (cash + GCash, not savings) at the end of that day or week if everything goes to plan. It jumps up on paydays and drops on bill days.",
  "Compare it with your real cash + GCash once in a while to see if you're on track.")
g("Actual spent",
  "YOU fill this in: what you really spent that day on living costs (fare, water, food, etc.), NOT bills. Leave it blank if you don't track that day.")
g("Under / (over)",
  "Planned living cost minus what you actually spent. A positive number = you spent LESS than planned (good). A number in brackets or red = you spent MORE than planned.")
g("Running under / (over)", "The total of all your Under / (over) so far. Positive = ahead of plan overall. Negative = behind.")
g("Week starts (Mon)|Week ends (Sun)", "First and last day of that week.")
g("Net change",
  "How much your spending money went up or down that week: Income − Bills − Living costs − Moved to savings. A negative week is normal when there's no payday; what matters is that the balance never runs out.")
g("Month|From|To", "The month, and the dates of it that are inside the plan (Oct 2026 is only Oct 27-31).")
g("Rent", "Your 2,000 monthly rent, due on the 1st. It includes electricity and water.")
g("SPayLater", "Your Shopee SPayLater installments (Nov 5 and Dec 5, 2026).")
g("Gloan", "Your three GCash GLoan installments for that month.")
g("Brother", "Your younger brother's allowance. 0 unless you switch it on (Assumptions row 35) or type an amount on Schedule.")
g("Other bills", "Bills that aren't rent, SPayLater, Gloan or your brother, e.g. the Smart data reload, or anything you add to Schedule.")
g("Net change in spendable",
  "How much your spending money went up (+) or down (−) over the month: Income − Total bills − Living total − Moved to savings. It does NOT count money moved to savings as lost: that money just moved to your savings pile.",
  "Example: a −1,700 month where you also saved 4,900 is still a good month.")
g("Spendable balance",
  "Money you can use for bills and daily costs (your cash + GCash), NOT counting savings. Shown at the last day of the month.",
  "It looks high at month-end because the month-end payday just arrived, and that money is needed for rent on the 1st and the early-month bills.")
g("Total money", "Spendable balance + Savings balance. Everything you have.")
g("Bills % of income",
  "The share of that month's salary eaten by bills. Nov-Dec 2026 is about 60%, so it's tight. Once the loans end it drops to about 17-20%.",
  "A common rule of thumb is to keep fixed bills under 50% of income.")
g("Type", "Income = money coming in. Bill = money going out on a fixed date.")
g("Category", "Groups bills for the Monthly sheet: Rent, SPayLater, Gloan, Family, Data, Salary, etc.")
g("Item", "What the payment is.")
g("Amount", "How much. Black = formula (e.g. linked to Assumptions), blue = a number you can type over.")
g("Paid?", "Type Yes once it's paid. The row turns grey with a line through it, so you can see what's done.")
g("Note", "Extra details about the row.")

# ---- hover comments on header cells ----
def norm(v):
    return str(v).replace(" (₱)", "").replace("\n", " ").strip()
def add_comments(ws, row):
    for c in ws[row]:
        if c.value is None: continue
        k = norm(c.value)
        if k in G:
            text, tip = G[k]
            body = text + ("\n\nTip: " + tip if tip else "")
            cm = Comment(body, "Budget guide"); cm.width = 320; cm.height = 40 + 15 * (len(body) // 45 + 1)
            c.comment = cm
for w_, r_ in ((wsP, 5), (wsD, 5), (wsW, 5), (wsM, 5), (wsC, 4)):
    add_comments(w_, r_)

# ---- Glossary sheet ----
wsG = wb.create_sheet("Glossary", 1)
ws = wsG
title(ws, "Glossary: what every heading means", "Plain-language explanations. Tip: hover over (or tap) any column heading in the other sheets to see the same note.")
widths(ws, {"A": 30, "B": 80, "C": 52})

r = 4
sec(ws, r, "How your money flows (read this first)"); r += 1
flow = [
    ("1. Payday comes in", "Every 15th and last day of the month you get about 6,000 (Pay in)."),
    ("2. Bills first", "Pay the fixed bills due before the next payday: rent, SPayLater, Gloan (Bills due)."),
    ("3. Pay yourself", "Move the planned amount to a separate savings account right away (Savings set-aside). Don't wait for 'leftovers'."),
    ("4. Live on the rest", "Fares, water, meals, misc come out of what's left (Living costs)."),
    ("5. Keep a cushion", "Whatever is left the day before the next payday is 'Money at period end'. Keep it above 500 (the safety buffer) for emergencies like a sudden fare or medicine."),
    ("6. Repeat", "That leftover carries into the next pay period as 'Money at start'."),
]
for a, b in flow:
    ws.cell(row=r, column=1, value=a).font = F(True)
    x = ws.cell(row=r, column=2, value=b); x.font = F(); x.alignment = WRAP
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 1
r += 1
sec(ws, r, "Live example from your plan: Oct 31, 2026 payday"); r += 1
ex = [
    ("Money at start", "='Pay Periods'!G7", "What you'll have left on Oct 30, before the payday"),
    ("+ Pay in", "='Pay Periods'!E7", "Oct 31 salary"),
    ("− Bills due", "='Pay Periods'!F7", "Rent Nov 1 + SPayLater Nov 5 + Gloan Nov 6"),
    ("− Savings set-aside", "='Pay Periods'!H7", "Nothing yet: the bills are too heavy this month"),
    ("− Living costs", "='Pay Periods'!K7", "Fares, water, misc for Oct 31 to Nov 14"),
    ("→ Money at period end", "='Pay Periods'!L7", "Left on Nov 14, the day before the next payday"),
    ("Cushion above buffer", "='Pay Periods'!M7", "Spare money above your 500 safety buffer"),
]
for a, f_, note in ex:
    ws.cell(row=r, column=1, value=a).font = F(a.startswith("→"))
    c = ws.cell(row=r, column=2, value=f_); c.number_format = PESO2; c.font = F(a.startswith("→"), GREEN); c.alignment = Alignment(horizontal="left")
    ws.cell(row=r, column=3, value=note).font = GREY9()
    for col in (1, 2, 3): ws.cell(row=r, column=col).border = BOX
    r += 1
r += 1
sec(ws, r, "Every heading, A to Z"); r += 1
header(ws, r, ["Heading", "What it means", "Tip / example"], height=22); r += 1
seen = set()
for k in sorted(G, key=lambda s: s.lower().lstrip("#")):
    text, tip = G[k]
    if k in ("#", "Day", "Ends", "From", "To", "Month", "Item", "Note", "Week ends (Sun)"): continue
    ws.cell(row=r, column=1, value=k).font = F(True)
    x = ws.cell(row=r, column=2, value=text); x.font = F(); x.alignment = WRAP
    y = ws.cell(row=r, column=3, value=tip); y.font = GREY9(); y.alignment = WRAP
    for col in (1, 2, 3): ws.cell(row=r, column=col).border = BOX
    ws.row_dimensions[r].height = max(15, 13 * (len(text) // 95 + 1))
    r += 1
r += 1
sec(ws, r, "Other words you'll see"); r += 1
other = [
    ("Safety buffer", "The minimum you never want to go below (500). It's for surprises, not for planned spending."),
    ("Emergency fund", "Savings kept only for real emergencies: getting sick, losing income, a broken phone you need for work. The usual target is 3 months of living costs + rent. Yours is set at 13,500."),
    ("13th month pay", "A yearly bonus required by Philippine law (1/12 of the basic salary you earned that year), paid on or before Dec 24. Counted as 0 here until you confirm it."),
    ("Handling fee", "An extra charge added by SPayLater on an installment."),
    ("Installment", "Paying a loan in smaller fixed amounts over several months (your Gloan and SPayLater)."),
    ("Spendable vs savings", "Spendable = money for bills and daily life. Savings = money you've set aside and don't touch. The plan keeps them apart so savings don't get spent by accident."),
    ("Numbers in red or with a minus", "Money going down or a shortfall. A dash (-) means zero."),
    ("Blue on yellow cells", "Inputs you can change. Black = calculated. Green = pulled from another sheet. Don't type over black cells."),
]
for a, b in other:
    ws.cell(row=r, column=1, value=a).font = F(True)
    x = ws.cell(row=r, column=2, value=b); x.font = F(); x.alignment = WRAP
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    for col in (1, 2, 3): ws.cell(row=r, column=col).border = BOX
    ws.row_dimensions[r].height = max(15, 13 * (len(b) // 120 + 1))
    r += 1
r += 1
sec(ws, r, "A few money habits that help"); r += 1
habits = [
    "Save on payday, not at the end. Move the Savings set-aside the same day your salary lands.",
    "Keep savings in a different account from your everyday GCash, so you don't spend it by accident.",
    "Avoid new SPayLater / GLoan while you're paying the current ones. Each new loan eats future paydays.",
    "Build the emergency fund before buying wants. Once it's full, you can save for goals (a laptop, travel, etc.).",
    "Check once a week: is your real cash + GCash close to the 'Planned balance' on Daily? If it's lower, find out where it went.",
]
for t in habits:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    x = ws.cell(row=r, column=1, value="• " + t); x.font = F(); x.alignment = WRAP
    r += 1
ws.freeze_panes = "A4"

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
    w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
    for rw in w.iter_rows():
        for c in rw:
            if c.font.name != FONT:
                nf = copy(c.font); nf.name = FONT
                if c.font.sz == 11: nf.sz = 10
                c.font = nf
for w, col in ((wsS, NAVY), (wsA, "BF9000"), (wsK, "BF9000"), (wsP, "BF9000"), (wsM, MID), (wsW, MID), (wsD, MID), (wsC, "7F7F7F"), (wsG, "548235")):
    w.sheet_properties.tabColor = col
wsD.print_title_rows = "5:5"; wsW.print_title_rows = "5:5"
# roll-up sheets have no inputs: protect them (no password) so formulas can't be deleted by accident
for w in (wsS, wsG, wsW, wsM):
    w.protection.sheet = True
    w.protection.formatColumns = False; w.protection.formatRows = False; w.protection.formatCells = False
    w.protection.autoFilter = False; w.protection.sort = False
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)

# python-side daily ledger for verification
ledger = []
bal = CASH + GCASH + BANK; sav = 0
for i in range(n):
    a, b = starts[i], ends[i]
    for k in range((b - a).days + 1):
        d = a + dt.timedelta(k)
        inc = sum(ev_amt(e) for e in events if e[0] == d and e[2] == "Income")
        bill = sum(ev_amt(e) for e in events if e[0] == d and e[2] == "Bill")
        c = day_cost(d); liv = c["total"] + (extras[i] if d == a else 0)
        s = savings_def[i] if d == a else 0
        bal = bal + inc - bill - liv - s; sav += s
        ledger.append((d.isoformat(), c["rides"], round(liv, 2), round(bal, 2), sav))
json.dump({"sim": sim, "ledger": ledger, "D": [D0, DL], "P": [P0, PL], "M": [M0r, MLr], "W": [W0, WL]}, open(os.path.join(HERE, "output", "ledger.json"), "w"))
print("saved", OUT, "events", len(events), "periods", n, "days", ndays, "weeks", nweeks)
for r in sim: print(r)
